import io
from unittest.mock import MagicMock
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.database import get_supabase_client
from app.main import app
from app.services.admin_import_service import AdminImportService
from app.utils.security import create_access_token

client = TestClient(app)

TEST_ADMIN_ID = str(uuid4())
TEST_EMAIL = "admin@natbell.com"


def _setup_admin_override():
    mock_db = MagicMock()
    admin_select = MagicMock()
    admin_eq = MagicMock()
    admin_eq.execute.return_value = MagicMock(
        data=[
            {
                "id": TEST_ADMIN_ID,
                "email": TEST_EMAIL,
                "name": "Admin Principal",
                "created_at": "2026-09-16T12:00:00Z",
            }
        ]
    )

    def table_router(name):
        m = MagicMock()
        if name == "admin_users":
            m.select.return_value = admin_select
            admin_select.eq.return_value = admin_eq
        return m

    mock_db.table.side_effect = table_router
    return mock_db


def test_generate_template_xlsx():
    content = AdminImportService.generate_template("xlsx")
    assert isinstance(content, bytes)
    assert len(content) > 0

    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    assert ws is not None
    assert ws.title == "Productos Natbell"
    headers = [cell.value for cell in ws[1]]
    assert "Nombre del Producto *" in headers
    assert "Precio Base *" in headers
    assert "SKU Variante" in headers
    assert ws.max_row >= 4  # Header + 3 sample rows


def test_generate_template_csv():
    content = AdminImportService.generate_template("csv")
    assert isinstance(content, bytes)
    text = content.decode("utf-8-sig")
    lines = text.strip().splitlines()
    assert len(lines) >= 4
    headers = lines[0].split(",")
    assert "nombre" in headers
    assert "precio_base" in headers
    assert "sku_variante" in headers


def test_download_template_unauthorized():
    res = client.get("/api/admin/import/template")
    assert res.status_code == 401


def test_download_template_authorized():
    mock_db = _setup_admin_override()
    token = create_access_token(subject=TEST_ADMIN_ID)
    app.dependency_overrides[get_supabase_client] = lambda: mock_db

    try:
        # XLSX
        res_xlsx = client.get(
            "/api/admin/import/template?format=xlsx",
            headers={"Authorization": f"Bearer {token}"},
        )
        assert res_xlsx.status_code == 200
        assert "spreadsheetml.sheet" in res_xlsx.headers["content-type"]
        assert "plantilla_productos_natbell.xlsx" in res_xlsx.headers["content-disposition"]

        # CSV
        res_csv = client.get(
            "/api/admin/import/template?format=csv",
            headers={"Authorization": f"Bearer {token}"},
        )
        assert res_csv.status_code == 200
        assert "text/csv" in res_csv.headers["content-type"]
        assert "plantilla_productos_natbell.csv" in res_csv.headers["content-disposition"]
    finally:
        app.dependency_overrides.pop(get_supabase_client, None)


def test_import_unauthorized():
    res = client.post(
        "/api/admin/import",
        files={"file": ("test.csv", b"nombre,precio_base\nTest,100", "text/csv")},
    )
    assert res.status_code == 401


def test_import_invalid_extension():
    mock_db = _setup_admin_override()
    token = create_access_token(subject=TEST_ADMIN_ID)
    app.dependency_overrides[get_supabase_client] = lambda: mock_db

    try:
        res = client.post(
            "/api/admin/import",
            files={"file": ("invalid.pdf", b"fake content", "application/pdf")},
            headers={"Authorization": f"Bearer {token}"},
        )
        assert res.status_code == 400
        assert "Formato no permitido" in res.json()["detail"]
    finally:
        app.dependency_overrides.pop(get_supabase_client, None)


def test_import_from_xlsx_success():
    mock_db = MagicMock()
    token = create_access_token(subject=TEST_ADMIN_ID)

    admin_select = MagicMock()
    admin_eq = MagicMock()
    admin_eq.execute.return_value = MagicMock(
        data=[{"id": TEST_ADMIN_ID, "email": TEST_EMAIL, "name": "Admin Principal"}]
    )

    cat_id = str(uuid4())
    subcat_id = str(uuid4())
    brand_id = str(uuid4())
    prod_id = str(uuid4())

    def table_router(name):
        m = MagicMock()
        if name == "admin_users":
            m.select.return_value = admin_select
            admin_select.eq.return_value = admin_eq
        elif name == "categories":
            m.select.return_value.execute.return_value = MagicMock(
                data=[{"id": cat_id, "name": "Cuidado Capilar", "slug": "cuidado-capilar"}]
            )
        elif name == "subcategories":
            m.select.return_value.execute.return_value = MagicMock(
                data=[{"id": subcat_id, "category_id": cat_id, "name": "Shampoos", "slug": "shampoos"}]
            )
        elif name == "brands":
            m.select.return_value.execute.return_value = MagicMock(
                data=[{"id": brand_id, "name": "Nov", "slug": "nov"}]
            )
        elif name == "products":
            # existing check
            m.select.return_value.ilike.return_value.limit.return_value.execute.return_value = MagicMock(data=[])
            # slug uniqueness check
            m.select.return_value.eq.return_value.execute.return_value = MagicMock(data=[])
            # insert
            m.insert.return_value.execute.return_value = MagicMock(data=[{"id": prod_id}])
        elif name == "product_variants":
            # sku check
            m.select.return_value.eq.return_value.limit.return_value.execute.return_value = MagicMock(data=[])
            # insert
            m.insert.return_value.execute.return_value = MagicMock(data=[{"id": str(uuid4())}])
        return m

    mock_db.table.side_effect = table_router
    app.dependency_overrides[get_supabase_client] = lambda: mock_db

    # Create XLSX in-memory
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["nombre", "categoria", "marca", "precio_base", "sku_variante", "stock_variante"])
    ws.append(["Shampoo Reparador", "Cuidado Capilar", "Nov", 8500.0, "SHP-REP-01", 15])
    excel_bytes = io.BytesIO()
    wb.save(excel_bytes)
    excel_content = excel_bytes.getvalue()

    try:
        res = client.post(
            "/api/admin/import",
            files={
                "file": (
                    "productos.xlsx",
                    excel_content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"Authorization": f"Bearer {token}"},
        )
        assert res.status_code == 200
        data = res.json()
        assert data["success"] is True
        assert data["products_created"] == 1
        assert data["variants_created"] == 1
        assert data["total_rows_processed"] == 1
    finally:
        app.dependency_overrides.pop(get_supabase_client, None)


def test_import_from_csv_success():
    mock_db = MagicMock()
    token = create_access_token(subject=TEST_ADMIN_ID)

    admin_select = MagicMock()
    admin_eq = MagicMock()
    admin_eq.execute.return_value = MagicMock(
        data=[{"id": TEST_ADMIN_ID, "email": TEST_EMAIL, "name": "Admin Principal"}]
    )

    cat_id = str(uuid4())
    subcat_id = str(uuid4())
    brand_id = str(uuid4())
    prod_id = str(uuid4())

    def table_router(name):
        m = MagicMock()
        if name == "admin_users":
            m.select.return_value = admin_select
            admin_select.eq.return_value = admin_eq
        elif name == "categories":
            m.select.return_value.execute.return_value = MagicMock(
                data=[{"id": cat_id, "name": "Cuidado Capilar", "slug": "cuidado-capilar"}]
            )
        elif name == "subcategories":
            m.select.return_value.execute.return_value = MagicMock(
                data=[{"id": subcat_id, "category_id": cat_id, "name": "Shampoos", "slug": "shampoos"}]
            )
        elif name == "brands":
            m.select.return_value.execute.return_value = MagicMock(
                data=[{"id": brand_id, "name": "Nov", "slug": "nov"}]
            )
        elif name == "products":
            m.select.return_value.ilike.return_value.limit.return_value.execute.return_value = MagicMock(data=[])
            m.select.return_value.eq.return_value.execute.return_value = MagicMock(data=[])
            m.insert.return_value.execute.return_value = MagicMock(data=[{"id": prod_id}])
        elif name == "product_variants":
            m.select.return_value.eq.return_value.limit.return_value.execute.return_value = MagicMock(data=[])
            m.insert.return_value.execute.return_value = MagicMock(data=[{"id": str(uuid4())}])
        return m

    mock_db.table.side_effect = table_router
    app.dependency_overrides[get_supabase_client] = lambda: mock_db

    csv_content = (
        "nombre,categoria,marca,precio_base,sku_variante,stock_variante\n"
        "Acondicionador Brillo,Cuidado Capilar,Nov,7900,ACON-BRILLO,20\n"
    ).encode("utf-8")

    try:
        res = client.post(
            "/api/admin/import",
            files={"file": ("productos.csv", csv_content, "text/csv")},
            headers={"Authorization": f"Bearer {token}"},
        )
        assert res.status_code == 200
        data = res.json()
        assert data["success"] is True
        assert data["products_created"] == 1
        assert data["variants_created"] == 1
    finally:
        app.dependency_overrides.pop(get_supabase_client, None)


def test_fuzzy_matching_taxonomies():
    mock_db = MagicMock()
    service = AdminImportService(mock_db)

    cat_id = str(uuid4())
    subcat_id = str(uuid4())
    brand_id = str(uuid4())

    taxonomies = {
        "categories": [{"id": cat_id, "name": "Cuidado Capilar", "slug": "cuidado-capilar"}],
        "subcategories": [{"id": subcat_id, "category_id": cat_id, "name": "Shampoos", "slug": "shampoos"}],
        "brands": [{"id": brand_id, "name": "L'Oréal", "slug": "loreal"}],
    }

    # Brand match: "Loreal" without apostrophe or accent
    resolved_brand = service._resolve_brand_id("Loreal", taxonomies)
    assert resolved_brand == brand_id

    # Subcategory match: "Shampoo" singular matches "Shampoos"
    resolved_subcat = service._resolve_subcategory_id("Cuidado Capilar", "Shampoo", taxonomies)
    assert resolved_subcat == subcat_id

    # Category match without subcat: "cuidado capilar" matches
    resolved_subcat2 = service._resolve_subcategory_id("cuidado capilar", "", taxonomies)
    assert resolved_subcat2 == subcat_id
